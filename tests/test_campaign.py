"""Step 4: plan -> specs -> build -> review sheet -> approvals -> library; compliance and spec validation."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import run
from lrstudio import compliance, library, spec, store, templates
from lrstudio.util import read_json, write_json

MCMINN_OLD = Path("/home/user/lawrank-creative-studio/engine/template/clients/mcminn")


@pytest.fixture
def client(home):
    if not MCMINN_OLD.exists():
        pytest.skip("needs the internal test client")
    out = run("client", "import", str(MCMINN_OLD))
    assert out["ok"], out
    return store.find("mcminn")


def good_option(key="A", **over):
    o = {"key": key, "template": "statement", "treatment": "dark",
         "copy": {"label": "Meet the firm", "headline": "A family-owned firm that puts clients first.", "accent": "puts clients first.",
                  "support": "Injury attorneys for Austin, Central Texas and Salt Lake City.", "cta": "Learn about our firm"},
         "facts": ["family_owned", "serves"], "caption": "McMinn Law Firm is family-owned and puts clients first.",
         "alt": "Dark graphic with the headline: A family-owned firm that puts clients first."}
    o.update(over)
    return o


def write_post(c, pid, options, campaign=None):
    from lrstudio import campaign as camp
    d = camp.campaign_dir(c, campaign)
    f = d / "posts" / f"{pid}.json"
    p = read_json(f)
    p["options"] = options
    write_json(f, p)
    return d


def test_campaign_new_follows_the_plan(client):
    out = run("campaign", "new", "mcminn", "--start", "2026-11")
    assert out["ok"] and out["posts"] == 30 and out["motion_posts"] == 0      # growth: all-package + growth posts
    assert read_json(Path(out["folder"]) / "posts" / "KIT.json")["format"] == "kit"
    d = Path(out["folder"])
    p = read_json(d / "posts" / "PI-003.json")
    assert p["format"] == "carousel" and p["plan"]["angles"] and set(p["suggested_templates"]) <= set(templates.catalog())
    run("client", "set", "mcminn", "package=authority")
    out = run("campaign", "new", "mcminn", "--name", "Authority run")
    assert out["posts"] == 48 and out["motion_posts"] == 12 and read_json(Path(out["folder"]) / "posts" / "PI-006.json")["format"] == "motion"


def test_build_review_approve_library(client):
    run("campaign", "new", "mcminn")
    d = write_post(client, "PI-001", [good_option("A"),
                                      good_option("B", template="photo-split", photo="team",
                                                  copy={"label": "Welcome", "headline": "Helping the injured get justice.",
                                                        "accent": "get justice.", "support": "Offices in Austin and Salt Lake City.",
                                                        "cta": "Meet our attorneys"}, facts=["tagline", "offices"],
                                                  caption="Helping the injured get justice, from Austin and Salt Lake City.")])
    out = run("build", "mcminn", "--post", "PI-001")
    assert out["ok"], out
    assert out["built"] == 2 and (d / "index.html").exists() and (d / "images" / "PI-001-A-4x5.png").exists()
    # a reviewer approves option B in Google Sheets
    wb = load_workbook(d / "review.xlsx"); ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "B":
            row[11].value = "Approved"; row[12].value = "Love it"
    wb.save(d / "review.xlsx")
    out = run("build", "mcminn", "--post", "PI-001")
    assert out["redrawn"] == 0 and out["status"]["approved_options"] == 1
    lib = library.load()
    assert len(lib) == 1 and lib[0]["option"] == "B" and lib[0]["client"] == "McMinn Law Firm"
    found = run("library", "search", "injured justice Austin")
    assert found["results"] and found["results"][0]["headline"].startswith("Helping")
    # a reviewer asks for changes on A: Claude sees it in `client show` and in the next build
    wb = load_workbook(d / "review.xlsx"); ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "A":
            row[11].value = "Changes requested"; row[12].value = "Warmer headline"
    wb.save(d / "review.xlsx")
    shown = run("client", "show", "mcminn")
    assert shown["campaigns"][-1]["feedback"] == {"PI-001 A": "Changes requested: Warmer headline"}
    assert run("build", "mcminn", "--post", "PI-001")["reviewer_feedback"] == {"PI-001 A": "Changes requested: Warmer headline"}
    # changing an approved option sends it back for review
    p = read_json(d / "posts" / "PI-001.json"); p["options"][1]["copy"]["support"] = "Offices in Austin and Utah."
    write_json(d / "posts" / "PI-001.json", p)
    run("build", "mcminn", "--post", "PI-001")
    ws = load_workbook(d / "review.xlsx").active
    st = {r[4].value: (r[11].value, r[12].value) for r in ws.iter_rows(min_row=2)}
    assert st["B"] == ("Changed after approval", "Love it")
    assert library.load() == []


def test_build_refuses_bad_specs_with_specific_fixes(client):
    run("campaign", "new", "mcminn")
    write_post(client, "PI-001", [good_option("A", copy=dict(good_option()["copy"], headline="x" * 90)),
                                  good_option("B", facts=["nope"], photo="ghost")])
    out = run("build", "mcminn", "--post", "PI-001")
    errs = " ".join(out["not_built"]["PI-001"])
    assert not out["ok"] and "copy.headline is 90 characters; max 64" in errs and "'nope'" in errs and "ghost" in errs


def test_unwritten_posts_are_reported(client):
    run("campaign", "new", "mcminn")
    out = run("build", "mcminn", "--post", "PI-001,PI-002")
    assert not out["ok"] and "PI-001" in out["todo"]
    out = run("build", "mcminn", "--month", "1")             # the kit post comes pre-written; the rest are listed
    assert out["built"] >= 7 and "PI-001" in out["todo"]


def test_repeats_are_flagged(client):
    run("campaign", "new", "mcminn")
    write_post(client, "PI-001", [good_option("A")])
    write_post(client, "PI-002", [good_option("A")])
    out = run("build", "mcminn", "--post", "PI-001,PI-002")
    assert any("reads almost the same" in r for r in out["repeats"])


def test_directions_without_people_photos(home, monkeypatch, tmp_path):
    import fixture_site
    monkeypatch.setenv("LR_ALLOW_PRIVATE", "1")
    url, stop = fixture_site.serve(fixture_site.build(tmp_path / "site"))
    try:
        run("client", "new", url)
        out = run("directions", "harbor")
    finally:
        stop()
    assert "A" not in out["directions"] and "skipped" in out["note"]
    assert out["ok"] and not out["not_built"] and out["built"] >= 8 and not out["flagged"], out


# ---------------------------------------------------------------- compliance + validation units
FACTS = {"founded": {"text": "Founded in 2005.", "category": "history", "usable": True, "status": "usable"},
         "verdict": {"text": "$4.2 million verdict for a crash victim.", "category": "results", "usable": False, "status": "needs approval"},
         "fees": {"text": "No fee unless we win.", "category": "fees", "usable": True, "status": "usable"},
         "review": {"text": "They answered every call.", "category": "testimonial", "usable": True, "status": "usable"}}
META = {"compliance": None}


def lint(copy, facts=(), caption="x", **kw):
    return compliance.lint({"copy": copy, "facts": list(facts), "caption": caption, **kw}, FACTS, META)


def levels(issues):
    return {(i["rule"], i["level"]) for i in issues}


def test_guarantees_and_superlatives_block():
    assert ("guarantee", "block") in levels(lint({"headline": "We guarantee results."}))
    assert ("superlative", "block") in levels(lint({"headline": "The best lawyers in Tampa."}))
    assert ("expert", "warn") in levels(lint({"headline": "Truck accident experts."}))


def test_numbers_must_come_from_cited_facts():
    assert ("unverified", "block") in levels(lint({"headline": "Serving you since 1998."}))
    assert ("unverified", "block") not in levels(lint({"headline": "Serving you since 2005."}, ["founded"]))
    assert ("unverified", "block") not in levels(lint({"headline": "3 things to do after a crash."}))


def test_results_need_a_fact_approval_and_the_disclaimer():
    lv = levels(lint({"headline": "$4.2 million verdict."}))
    assert ("results", "block") in lv and ("disclaimer", "block") in lv
    lv = levels(lint({"headline": "$4.2 million verdict.", "disclaimer": compliance.DISCLAIMER}, ["verdict"]))
    assert ("disclaimer", "block") not in lv and ("approval", "warn") in lv


def test_fee_promises_and_testimonials():
    assert ("fees", "warn") in levels(lint({"headline": "No fee unless we win."}))
    assert ("fees", "warn") not in levels(lint({"headline": "No fee unless we win."}, ["fees"]))
    assert ("testimonial", "block") in levels(compliance.lint({"copy": {"quote": "Great firm"}, "facts": []}, FACTS, {"compliance": "testimonial"}))


def test_slot_rules_for_carousels():
    errs = []
    meta = templates.get("photo-story")
    spec.check_slots({"cover": {"headline": "ok"}, "slides": [{"title": "a" * 50}], "close": {"headline": "x"}}, meta["slots"], "copy", errs)
    joined = " ".join(errs)
    assert "copy.slides needs at least 3" in joined and "copy.slides[0].title is 50 characters" in joined


def test_educational_words_are_not_results():
    assert not any(i["rule"] in ("results", "disclaimer") for i in lint({"answer": "Many claims are resolved through a settlement; some reach a jury verdict."}))
    hit = [i for i in lint({"headline": "Our team recovered a settlement for a family."}) if i["rule"] == "results"]
    assert hit and "copy.headline" in hit[0]["msg"] and "recovered" in hit[0]["msg"]


def test_quoting_a_usable_fact_with_free_is_not_a_fee_warning():
    facts = dict(FACTS, contact={"text": "Call or use the free case evaluation form.", "category": "contact", "usable": True, "status": "usable"})
    iss = compliance.lint({"copy": {"support": "Use the free case evaluation form."}, "facts": ["contact"], "caption": "x"}, facts, META)
    assert ("fees", "warn") not in levels(iss)


def test_confirm_lines_reach_the_reviewer(client):
    run("campaign", "new", "mcminn")
    d = write_post(client, "PI-001", [good_option("A", confirm=["The firm's founding year"])])
    out = run("build", "mcminn", "--post", "PI-001")
    assert "The firm's founding year" in out["check_before_posting"]["PI-001 A"]
    ws = load_workbook(d / "review.xlsx").active
    assert any("founding year" in str(r[10].value) for r in ws.iter_rows(min_row=2))


def test_package_upgrade_adds_posts_and_motion(client):
    out = run("campaign", "new", "mcminn")
    write_post(client, "PI-001", [good_option("A")])
    run("client", "set", "mcminn", "package=authority")
    up = run("campaign", "update", "mcminn")
    assert len(up["added"]) == 18 and "PI-006" in up["added"] and "PI-006" in up["motion_posts"]
    assert read_json(Path(out["folder"]) / "posts" / "PI-006.json")["format"] == "motion"
    assert read_json(Path(out["folder"]) / "posts" / "PI-001.json")["options"]          # written posts untouched


def test_confirm_as_a_string_and_revised_status(client):
    run("campaign", "new", "mcminn")
    d = write_post(client, "PI-001", [good_option("A", confirm="The founding year")])
    out = run("build", "mcminn", "--post", "PI-001")
    assert out["check_before_posting"]["PI-001 A"] == ["The founding year"]
    wb = load_workbook(d / "review.xlsx"); ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "PI-001":
            row[11].value = "Changes requested"; row[12].value = "Shorter"
    wb.save(d / "review.xlsx")
    p = read_json(d / "posts" / "PI-001.json"); p["options"][0]["copy"]["support"] = "Injury attorneys for Austin and Utah."
    write_json(d / "posts" / "PI-001.json", p)
    run("build", "mcminn", "--post", "PI-001")
    ws = load_workbook(d / "review.xlsx").active
    assert [r[11].value for r in ws.iter_rows(min_row=2) if r[0].value == "PI-001"] == ["Revised"]
