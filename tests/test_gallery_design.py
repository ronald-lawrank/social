"""Regressions from onboarding Ben Carrasco (Austin family law, October 2026): the gallery's build record, the Design
tab, the brand board's sample copy and the short name. Drawing is faked, so none of these need Chromium."""
import json
import re
from pathlib import Path

import pytest

from conftest import run
from lrstudio import board, build, render, store, templates
from lrstudio.util import read_json, write_json

BRAND = {"name": "The Law Office of Ben Carrasco, PLLC", "short": "Ben Carrasco", "website": "bencarrascolaw.com",
         "colors": {"ink": "#0B253E", "accent": "#1C5373", "paper": "#F5F5F5", "muted": "#5F6772", "body": "#404040",
                    "line": "#D9DCDF", "accentBright": "#2B7FB1"},
         "fonts": {"head": {"family": "Jost", "weight": 700}, "body": {"family": "Jost", "weight": 400},
                   "ui": {"family": "Figtree", "weight": 400}},
         "cta": {"radius": 0, "upper": True}, "treatments": {}, "logo": {"status": "raster"},
         "voice": {"tagline": "Facing Divorce is TOUGH", "positioning": None, "notes": []},
         "practice_areas": ["Divorce", "Family law"], "offices": [{"city": "Austin", "state": "TX"}], "people": []}
FACTS = [{"key": "tagline", "text": "Facing Divorce is TOUGH", "category": "messaging", "verified": True},
         {"key": "family_owned", "text": "We value family-owned businesses in a divorce.", "category": "differentiator", "verified": True},
         {"key": "fees", "text": "Call to schedule a risk-free consultation.", "category": "fees", "verified": True, "needs_approval": True},
         {"key": "awards", "text": "Board Certified in Family Law.", "category": "awards", "verified": True, "needs_approval": True}]


def statement(key, headline):
    return {"key": key, "template": "statement", "treatment": "accent", "facts": [],
            "copy": {"label": "Ask a lawyer", "headline": headline}, "caption": headline, "alt": f"Graphic: {headline}"}


@pytest.fixture
def firm(home):
    c = store.create(BRAND["name"], "https://bencarrascolaw.com", practice_area="Family law", direction="D")
    write_json(c.brand_dir / "brand.json", BRAND)
    write_json(c.brand_dir / "photos" / "photos.json", {"ben-2": {"file": "ben-2.jpg", "kind": "person", "name": "Ben Carrasco",
                                                                 "use": True, "w": 768, "h": 512, "faces": []}})
    store.write_facts(c, FACTS)
    d = c.campaigns_dir / "2026-10 Six-Month"
    posts = {"PI-001": (1, [statement("A", "Facing divorce? Start here."), statement("B", "Know your options.")]),
             "PI-002": (1, [statement("A", "Custody questions, answered.")]),
             "PI-009": (2, [statement("A", "What to gather before you file.")]),
             "KIT": (1, [{"key": "FB", "template": "facebook-cover", "copy": {"headline": "Austin family law"}, "facts": [], "photo": "ben-2",
                          "caption": "Cover", "alt": "Facebook cover"},
                         {"key": "LI", "template": "linkedin-banner", "copy": {"headline": "Austin family law"}, "facts": [], "photo": "ben-2",
                          "caption": "Banner", "alt": "LinkedIn banner"}])}
    for pid, (month, opts) in posts.items():
        write_json(d / "posts" / f"{pid}.json", {"id": pid, "month": month, "week": 1, "topic": pid, "slot_type": "Firm Introduction",
                                                 "format": "kit" if pid == "KIT" else "static", "options": opts})
    write_json(d / "campaign.json", {"name": d.name})
    return c


@pytest.fixture
def drawn(monkeypatch):
    """Fake renderer: honours the cache like the real one and records what it actually drew."""
    log = []

    def fake(client_dir, brand, designs, out_dir, cache=None, flat=False):
        out = []
        for x in designs:
            if (cache or {}).get(x["id"]):
                out.append(dict(cache[x["id"]], cached=True)); continue
            log.append(x["id"])
            out.append({"id": x["id"], "key": "k", "template": x["template"], "size": x["size"], "html": f"previews/{x['id']}.html",
                        "images": [f"images/{x['id']}.png"], "photos": [], "issues": [], "frames": 1})
        return out
    monkeypatch.setattr(render, "render", fake)
    return log


def test_partial_and_forced_builds_merge_into_the_build_record(firm, drawn):
    d = firm.campaigns_dir / "2026-10 Six-Month"
    out = build.build(firm, month=1)
    assert out["not_built"] == {} and out["built"] == 5, out
    every = set(read_json(d / ".build.json"))
    assert {"KIT-FB-fb-cover", "KIT-LI-li-banner", "PI-001-A-4x5", "PI-002-A-4x5"} <= every
    drawn.clear()
    build.build(firm, posts="PI-001", force=True)                    # --force redraws only the selected post ...
    assert sorted(drawn) == ["PI-001-A-4x5", "PI-001-B-4x5"]
    assert set(read_json(d / ".build.json")) == every                # ... and the rest of the record survives
    build.build(firm, month=2)                                       # a later month keeps the kit (dashed sizes)
    after = set(read_json(d / ".build.json"))
    assert every | {"PI-009-A-4x5"} == after
    page = (d / "index.html").read_text()
    assert "not drawn yet" not in page and "KIT-FB-fb-cover.png" in page


def test_deleted_posts_leave_the_build_record(firm, drawn):
    d = firm.campaigns_dir / "2026-10 Six-Month"
    build.build(firm, month=1)
    (d / "posts" / "PI-002.json").unlink()
    build.build(firm, posts="PI-001")
    ids = set(read_json(d / ".build.json"))
    assert not any(k.startswith("PI-002-") for k in ids) and "KIT-FB-fb-cover" in ids
    assert build.post_of("KIT-FB-fb-cover", {"KIT", "PI-001"}) == "KIT"
    assert build.post_of("DIR-A-A1-4x5", {"DIR-A", "DIR-B"}) == "DIR-A"


def test_gallery_has_a_design_tab_built_from_the_brand_and_facts(firm, drawn):
    d = firm.campaigns_dir / "2026-10 Six-Month"
    from PIL import Image
    if not (firm.brand_dir / "photos" / "ben-2.jpg").exists():
        Image.new("RGB", (768, 512), "#777").save(firm.brand_dir / "photos" / "ben-2.jpg")
    build.build(firm, month=1)
    index = (d / "index.html").read_text()
    assert not (d / "design.html").exists()                                     # one file: the guide is a tab
    for tab in ("posts", "kit", "design"):
        assert f"data-t='{tab}'" in index and f"id='tab-{tab}'" in index
    posts, kit, design = (index.split(f"id='tab-{t}'")[1].split("<div class='tabp'")[0] for t in ("posts", "kit", "design"))
    assert "PI-001" in posts and "KIT" not in posts                             # the kit has its own tab
    assert "KIT-FB-fb-cover.png" in kit and "PI-001" not in kit
    assert "#tab-design section{" in index and "\nsection{" not in index       # the guide's CSS can't touch the gallery
    for section in ("Direction", "Logo", "Colour", "Typography", "Treatments", "Photography", "Formats", "Words"):
        assert f"<h2>{section}</h2>" in design
    assert "Brand colour" in design and "100% accent" in design                  # direction D and the campaign's mix
    assert "#1C5373" in design and re.search(r"White on ink</span><span class='r ok'>\d+\.\d:1", design)
    assert "Jost" in design and "Figtree" in design and "30 px" in design
    assert "Facebook cover" in design and "1640 × 624" in design and "2256 × 382" in design
    assert "Facing Divorce is TOUGH" in design                                  # the tagline, locked as written
    assert "Waiting for the firm's approval (2)" in design and "Board Certified in Family Law." in design
    assert "No &quot;free&quot;" in design and "Talk to us" in design and "Free consultation" not in design
    assert not re.search(r"injur|accident|justice|family-owned", design, re.I)  # no sample copy from another practice
    assert "Austin family law" in design
    assert "ben-2.jpg" in design and "Small original (768 px)" in design
    assert "../../brand/" not in index                                    # works on its own: shared, zipped or previewed
    assert (d / "previews" / "assets" / "brand" / "photos" / "ben-2.jpg").exists()
    assert "align-items:center" in index.split(".pair{")[1].split("}")[0]  # contrast pills sit in the row


def test_brand_board_copy_fits_the_practice_and_the_fees_approval(firm):
    bc = board.board_copy(firm)
    words = " ".join(bc["headline"] + bc["tile"]).lower()
    assert "divorce" in words or "custody" in words
    assert not re.search(r"injur|accident|justice|family-owned|puts clients first", words)
    assert bc["cta"] == "Talk to us"                                            # the fees fact isn't approved yet
    from openpyxl import load_workbook
    f = firm.brand_dir / "facts.xlsx"
    wb = load_workbook(f); ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "fees":
            row[7].value = "Yes"
    wb.save(f)
    assert board.board_copy(firm)["cta"] == "Free consultation"
    html = (Path(templates.get("brand-board")["dir"]) / "template.html").read_text()
    assert "Helping the injured" not in html and "family-owned" not in html and "cta('Free consultation')" not in html


@pytest.mark.parametrize("name,short", [
    ("The Law Office of Ben Carrasco, PLLC", "Ben Carrasco"), ("Law Offices of John Smith, P.C.", "John Smith"),
    ("The McMinn Law Firm", "McMinn"), ("McMinn Law Firm", "McMinn"), ("Smith & Jones Law Firm", "Smith"),
    ("The Law Firm of Jane Doe", "Jane Doe"), ("Ramos Injury Lawyers", "Ramos"), ("The Doe Firm", "Doe")])
def test_short_name_drops_articles_and_law_office_of(name, short):
    assert store.short_name(name) == short


def test_refresh_repairs_a_bare_article_short_name(firm, monkeypatch):
    monkeypatch.setattr(board, "render_board", lambda c: None)
    b = read_json(firm.brand_dir / "brand.json"); b["short"] = "The"
    write_json(firm.brand_dir / "brand.json", b)
    out = run("client", "refresh", firm.name)
    assert read_json(firm.brand_dir / "brand.json")["short"] == "Ben Carrasco"
    assert any("Ben Carrasco" in n for n in out["notes"])


def test_campaigns_use_a_practice_plan_when_there_is_one(firm, home):
    from lrstudio import campaign
    assert campaign.area_group("Divorce") == "family law" and campaign.area_group("Personal injury") == "personal injury"
    out = run("campaign", "new", firm.name, "--name", "Default plan")
    assert out["ok"] and "personal injury" in out["plan_note"] and "plan/family-law.csv" in out["plan_note"]
    src = campaign.plan_file()[0].read_text(encoding="utf-8-sig").splitlines()
    rows = [src[0]] + [l.replace("PI-0", "FL-0", 1) for l in src[1:4]]
    (home / "LawRank Studio" / "plan").mkdir()
    (home / "LawRank Studio" / "plan" / "family-law.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    out = run("campaign", "new", firm.name, "--name", "Family plan")
    posts = sorted(p.stem for p in (firm.campaigns_dir / "Family plan" / "posts").glob("*.json"))
    assert out["plan_file"] == "family-law.csv" and "plan_note" not in out
    assert posts[0].startswith("FL-0") and not any(p.startswith("PI-") for p in posts)


def test_motion_families_offered_for_the_firms_practice(firm):
    out = run("templates", "motion", "--client", firm.name)
    assert out["ok"] and [f["family"] for f in out["families"]] == ["flip", "scenes"] and "journey" in out["not_for_this_practice"]
    assert len(run("templates", "motion")["families"]) == 5                    # composed videos suit every practice


# ---- new-run fixes from the Alders & Lewellyn onboarding
def test_third_party_logos_are_never_picked(monkeypatch):
    from types import SimpleNamespace
    from lrstudio import onboard
    site = SimpleNamespace(url="https://aldersandlewellyn.com/", final_url="https://aldersandlewellyn.com/", logos=[
        {"url": "https://cdn.trustindex.io/assets/platform/Google/logo.svg", "where": "page", "alt": "", "w": 120},
        {"url": "https://aldersandlewellyn.com/wp-content/uploads/aldersandlewellyn-logo.png", "where": "header", "alt": "", "w": 400}])
    fetched = []

    def fake_download(url, referer=None):
        fetched.append(url)
        return (b"<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 10 10'/>", "image/svg+xml") if url.endswith(".svg") else (b"PNG", "image/png")
    monkeypatch.setattr(onboard, "download", fake_download)
    best = onboard.pick_logo(site)
    assert best["url"].endswith("aldersandlewellyn-logo.png") and not any("trustindex" in u for u in fetched)


def test_small_text_reads_on_a_light_accent():
    from lrstudio import theme
    from lrstudio.extract import contrast
    t = theme.treatments({"colors": {"ink": "#344952", "accent": "#C9A977", "paper": "#F4F4F4"}})["accent"]
    assert contrast(t["fg"], "#C9A977") >= 4.5


def test_kit_designs_are_read_right_when_packaging():
    from lrstudio import deliver
    posts = {"KIT", "PI-001", "DIR-A"}
    assert deliver.parse_id("KIT-FB-fb-cover", posts) == ("KIT", "FB")
    assert deliver.parse_id("KIT-LI-li-banner", posts) == ("KIT", "LI")
    assert deliver.parse_id("DIR-A-A1-4x5", posts) == ("DIR-A", "A1")
    assert deliver.parse_id("PI-001-A-4x5", posts) == ("PI-001", "A")


def test_test_designs_never_use_unapproved_or_invented_lines(home):
    from lrstudio import campaign as camp
    run("client", "new", "https://alders-example.example/", "--manual", "--name", "Alders & Lewellyn, PLLC")
    c = store.find("alders")
    b = c.brand(); b["voice"] = {"tagline": "Alders & Lewellyn: Memphis Personal Injury Attorneys"}
    write_json(c.brand_dir / "brand.json", b)
    run("client", "set", "alders", "practice_area=Personal injury")
    camp.directions(store.find("alders"))
    text = json.dumps([read_json(f) for f in (c.campaigns_dir / "Test designs" / "posts").glob("DIR-*.json")])
    for bad in ("Memphis Personal Injury Attorneys", "Our promise", "Straight answers", "Here when you need answers"):
        assert bad not in text
    assert "Do I need a lawyer after an accident?" in text


GOOGLE_BADGE = ("<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 266 70'>" + "<path d='M1 1h5v5z' fill='#FBBC05'/>" * 5
                + "<path d='M2 2h1v1z' fill='#4285F4'/><path d='M3 3h1v1z' fill='#34A853'/><path d='M4 4h1v1z' fill='#EA4335'/></svg>")


def test_review_badges_are_never_the_logo():
    from types import SimpleNamespace
    from lrstudio import onboard
    assert onboard.third_party_svg(GOOGLE_BADGE)
    firm = "<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 300 80'><path d='M0 0h300v80H0z' fill='#0096E3'/></svg>"
    assert not onboard.third_party_svg(firm)
    site = SimpleNamespace(url="https://jhartley.example/", final_url="https://jhartley.example/", logos=[
        {"svg": GOOGLE_BADGE, "where": "page"}, {"svg": firm, "where": "header"}])
    assert "#0096E3" in onboard.pick_logo(site)["svg"]


def test_banners_and_site_graphics_arent_used_as_photos(tmp_path, monkeypatch):
    import io
    import random
    from PIL import Image
    from lrstudio import crawl, photos

    def img(w, h):
        rnd = random.Random(w * h)
        im = Image.frombytes("RGB", (w, h), bytes(rnd.randrange(256) for _ in range(w * h * 3)))
        buf = io.BytesIO(); im.save(buf, "JPEG"); return buf.getvalue()
    files = {"https://f.example/DESKTOP-1.jpg": img(2400, 825), "https://f.example/Client-Review.png": img(1472, 900),
             "https://f.example/team-room.jpg": img(1600, 1100)}
    monkeypatch.setattr(crawl, "download", lambda url, referer=None: (files[url], "image/jpeg"))
    out = photos.collect([{"url": u, "page": "https://f.example/"} for u in files], tmp_path, [], "f.example")
    why = {m["image_url"].rsplit("/", 1)[1]: (m["use"], m["why"]) for m in out.values()}
    assert why["DESKTOP-1.jpg"][0] is False and "banner" in why["DESKTOP-1.jpg"][1]
    assert why["Client-Review.png"][0] is False and "site graphic" in why["Client-Review.png"][1]
    assert why["team-room.jpg"][0] is True, why


def test_seo_titles_arent_used_as_taglines():
    from lrstudio import campaign
    assert campaign.DESCRIPTOR.match("California Personal Injury Lawyer")
    assert not campaign.DESCRIPTOR.match("Helping the Injured Get Justice.")
    assert campaign._sentence("California Personal Injury Lawyer") == "California Personal Injury Lawyer."
