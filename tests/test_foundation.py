import json

import pytest

from conftest import run
from lrstudio import paths, store, util


def test_home_from_env(home):
    assert paths.home() == home


def test_home_search_finds_drive_layout(tmp_path, monkeypatch):
    monkeypatch.delenv("LR_HOME", raising=False)
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("LR_STATE_DIR", str(tmp_path / "state"))
    h = tmp_path / "Library/CloudStorage/GoogleDrive-a@lawrank.com/Shared drives/LawRank/Social Media (DNU)/Claude"
    (h / "LawRank Studio").mkdir(parents=True); (h / "Clients").mkdir()
    monkeypatch.setattr(paths.Path, "home", classmethod(lambda cls: tmp_path))
    assert paths.home() == h
    assert paths.settings()["home"] == str(h)


def test_missing_home_is_a_clear_error(tmp_path, monkeypatch):
    monkeypatch.delenv("LR_HOME", raising=False)
    monkeypatch.setenv("LR_STATE_DIR", str(tmp_path / "state"))
    monkeypatch.setattr(paths.Path, "home", classmethod(lambda cls: tmp_path))
    with pytest.raises(paths.HomeNotFound):
        paths.home()


def test_create_and_find_client(home):
    c = store.create("Smith & Jones Injury Lawyers", "www.SmithJones.com/")
    assert c.dir.name == "Smith & Jones Injury Lawyers" and c.data["domain"] == "smithjones.com"
    assert store.find("smith & jones injury lawyers").slug == c.slug
    assert store.find("smithjones.com").slug == c.slug
    assert store.find("https://www.smithjones.com/contact").slug == c.slug
    assert store.find("jones").slug == c.slug
    with pytest.raises(util.LRError):
        store.create("Another name", "smithjones.com")
    with pytest.raises(util.LRError) as e:
        store.find("nobody")
    assert "lr client" in e.value.fix


def test_folder_names_are_safe(home):
    c = store.create('Lee: "Trial" Lawyers / P.C.', "lee.test")
    assert c.dir.name == "Lee Trial Lawyers P.C"


def test_facts_merge_keeps_human_edits(home):
    from openpyxl import load_workbook
    c = store.create("Firm", "firm.test")
    store.write_facts(c, [
        {"key": "founded", "text": "Founded in 1998.", "category": "history", "source": "https://firm.test/about", "excerpt": "Founded in 1998.", "verified": True},
        {"key": "results", "text": "$4.2 million verdict.", "category": "results", "source": "https://firm.test/results", "excerpt": "$4.2 million", "verified": True, "needs_approval": True},
    ])
    f = store.read_facts(c)
    assert f["founded"]["usable"] and not f["results"]["usable"] and f["results"]["status"] == "needs approval"
    # a reviewer approves the result and edits the founded text in Google Sheets
    path = c.brand_dir / "facts.xlsx"
    wb = load_workbook(path); ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "results": row[7].value = "Yes"; row[8].value = "Dana (reviewer)"
        if row[0].value == "founded": row[1].value = "Serving Tampa since 1998."
    wb.save(path)
    # a rescan proposes the old text again plus a new fact
    store.write_facts(c, [{"key": "founded", "text": "Founded in 1998.", "verified": True},
                          {"key": "languages", "text": "Se habla español.", "verified": True}])
    f = store.read_facts(c)
    assert f["results"]["usable"] and f["results"]["approved_by"] == "Dana (reviewer)"
    assert f["founded"]["text"] == "Serving Tampa since 1998."
    assert "languages" in f


def test_lock_blocks_a_live_holder_and_takes_over_stale(tmp_path):
    with util.Lock(tmp_path, "build month 2"):
        info = json.loads((tmp_path / ".lr-lock").read_text())
        info["pid"] = 1  # pid 1 is alive on every unix box: simulates another live run
        (tmp_path / ".lr-lock").write_text(json.dumps(info))
        with pytest.raises(util.LRError) as e:
            with util.Lock(tmp_path, "build month 3"):
                pass
        assert "tester" in str(e.value) or "already running" in str(e.value)
    assert not (tmp_path / ".lr-lock").exists()
    (tmp_path / ".lr-lock").write_text(json.dumps({"user": "x", "host": "elsewhere", "pid": 5, "what": "w", "since": "2020-01-01T00:00:00"}))
    with util.Lock(tmp_path, "stale is replaced"):
        pass


def test_history_log(home):
    c = store.create("Firm", "firm.test")
    util.log(c.dir, "lr build month 2")
    assert "tester | lr build month 2 | ok" in (c.dir / "history.log").read_text()


def test_cli_outputs_one_json_line_and_errors_have_fixes(home):
    out = run("client", "show", "nobody")
    assert out["ok"] is False and out["fix"] and out["_code"] == 1
    out = run("client", "list")
    assert out["ok"] and out["clients"] == []


def test_atomic_write_leaves_no_temp_files(tmp_path):
    util.write_json(tmp_path / "a.json", {"x": 1})
    assert [p.name for p in tmp_path.iterdir()] == ["a.json"]
