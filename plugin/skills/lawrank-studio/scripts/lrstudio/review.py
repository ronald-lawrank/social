"""The review loop: review.xlsx (people set Status in Google Sheets), index.html (the gallery) and flagged.jpg (only the
designs the checks stopped, for Claude to look at)."""
import hashlib
import html
import io
import json
from pathlib import Path

from . import spec, store
from .util import read_json, write_bytes, write_json, write_text

STATUSES = ["Draft", "Approved", "Changes requested", "Rejected"]
COLS = ["Post", "Month", "Week", "Topic", "Option", "Template", "Headline", "Caption", "Image", "Checks",
        "Check before posting", "Status", "Notes", "Version"]


def version(option: dict) -> str:
    return hashlib.sha1(json.dumps(option, sort_keys=True).encode()).hexdigest()[:8]


def read_sheet(d: Path) -> dict:
    """{(post, option): {"status", "notes", "version"}} from review.xlsx (what reviewers typed)."""
    f = d / "review.xlsx"
    if not f.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(f, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "") for h in next(rows, [])]
    ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in rows:
        g = lambda k: (r[ix[k]] if k in ix and ix[k] < len(r) else None)
        if g("Post") and g("Option"):
            st = str(g("Status") or "Draft").strip()
            st = next((s for s in STATUSES if s.lower() == st.lower()), "Changes requested" if st.lower().startswith("change") else st or "Draft")
            out[(str(g("Post")), str(g("Option")))] = {"status": st, "notes": g("Notes") or "", "version": g("Version") or ""}
    return out


def feedback(d: Path, posts: set | None = None, sheet: dict | None = None) -> dict:
    """What reviewers asked for: {"PI-001 B": "Changes requested: warmer headline"} for options not yet approved."""
    sheet = read_sheet(d) if sheet is None else sheet
    out = {}
    for (pid, key), v in sorted(sheet.items()):
        if posts is not None and pid not in posts:
            continue
        if v["status"] in ("Changes requested", "Rejected") or (v["notes"] and v["status"] == "Draft"):
            out[f"{pid} {key}"] = f"{v['status']}: {v['notes']}".rstrip(": ")
    return out


def _summary_line(o: dict) -> str:
    c = o.get("copy") or {}
    for k in ("headline", "question", "quote", "stat_caption", "myth"):
        if c.get(k):
            return c[k]
    if isinstance(c.get("cover"), dict):
        return c["cover"].get("headline", "")
    return ""


def update(c: store.Client, d: Path, built_posts: list[dict], designs: list[dict], notes: dict, repeats: list[str]) -> list[dict]:
    """Rewrite review.xlsx for every written post in the campaign, keeping reviewers' Status and Notes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    prev = read_sheet(d)
    cache = read_json(d / ".build.json", {}) or {}
    facts, pics = store.read_facts(c), read_json(c.brand_dir / "photos" / "photos.json", {}) or {}
    rows = []
    for f in sorted((d / "posts").glob("*.json")):
        p = read_json(f)
        if not p or not p.get("options"):
            continue
        comp = notes.get(p["id"]) or spec.validate_post(p, facts, pics)["issues"]
        for o in p["options"]:
            ids = [k for k in cache if k.startswith(f"{p['id']}-{o['key']}-")]
            res = [cache[k] for k in sorted(ids)]
            blocking = [i["msg"] for r in res for i in r.get("issues", []) if i["level"] == "block"]
            warn = [i["msg"] for r in res for i in r.get("issues", []) if i["level"] == "warn"]
            confirm = [i["msg"] for i in comp.get(o["key"], []) if i["level"] == "warn"]
            rep = [r for r in repeats if r.startswith(f"{p['id']} {o['key']}")]
            v = version(o)
            old = prev.get((p["id"], o["key"]), {})
            status = old.get("status", "Draft")
            if status == "Approved" and old.get("version") and old["version"] != v:
                status = "Changed after approval"
            img = next((im for r in res for im in r.get("images", [])), None)
            rows.append({"post": p["id"], "month": p.get("month"), "week": p.get("week"), "topic": p.get("topic", ""),
                         "slot_type": p.get("slot_type"), "format": p.get("format"), "option": o["key"], "template": o["template"],
                         "headline": _summary_line(o), "caption": o.get("caption", ""), "image": img,
                         "images": [im for r in res for im in r.get("images", [])], "html": res[0]["html"] if res else None,
                         "blocking": blocking, "warnings": warn + rep, "confirm": confirm, "status": status,
                         "notes": old.get("notes", ""), "version": v, "option_data": o})
    wb = Workbook(); ws = wb.active; ws.title = "Review"
    ws.append(COLS)
    for r in rows:
        checks = "Needs fixing: " + "; ".join(r["blocking"]) if r["blocking"] else ("Look at: " + "; ".join(r["warnings"]) if r["warnings"] else "OK")
        ws.append([r["post"], r["month"], r["week"], r["topic"], r["option"], r["template"], r["headline"], r["caption"],
                   r["image"] or "", checks, "; ".join(r["confirm"]), r["status"], r["notes"], r["version"]])
    head = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = head
    for i, w in enumerate([10, 7, 6, 34, 8, 14, 40, 60, 28, 40, 40, 18, 40, 10]):
        ws.column_dimensions[chr(65 + i)].width = w
    fills = {"Approved": "D1FAE5", "Changes requested": "FEF3C7", "Rejected": "FEE2E2", "Changed after approval": "FDE68A"}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        st = row[11].value
        if st in fills:
            row[11].fill = PatternFill("solid", fgColor=fills[st])
        if str(row[9].value).startswith("Needs"):
            row[9].fill = PatternFill("solid", fgColor="FEE2E2")
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"L2:L{max(2, ws.max_row + 50)}")
    ws.freeze_panes = "C2"
    buf = io.BytesIO(); wb.save(buf)
    write_bytes(d / "review.xlsx", buf.getvalue())
    return rows


def status(d: Path, rows: list[dict]) -> dict:
    posts = {r["post"] for r in rows}
    total = len(list((d / "posts").glob("*.json")))
    approved_posts = {r["post"] for r in rows if r["status"] == "Approved"}
    s = {"posts": total, "written": len(posts), "options": len(rows), "approved_options": sum(r["status"] == "Approved" for r in rows),
         "posts_with_an_approved_option": len(approved_posts), "changes_requested": sum(r["status"] == "Changes requested" for r in rows),
         "needs_fixing": sum(bool(r["blocking"]) for r in rows)}
    write_json(d / "status.json", s)
    return s


def flagged_sheet(d: Path, flagged: list[dict]) -> Path:
    """Only the designs the checks stopped, with the problem areas outlined in red."""
    from PIL import Image, ImageDraw
    tiles = []
    for x in flagged[:24]:
        r = x["result"]
        for n, rel in enumerate(r.get("images", [])):
            issues = [i for i in r.get("issues", []) if i["level"] == "block" and i.get("frame", 0) == n]
            if not issues:
                continue
            im = Image.open(d / rel).convert("RGB")
            dr = ImageDraw.Draw(im)
            for i in issues:
                if i.get("box"):
                    bx, by, bw, bh = i["box"]
                    dr.rectangle([bx, by, bx + bw, by + bh], outline=(230, 0, 0), width=8)
            s = 520 / im.height
            tiles.append((f"{x['id']}" + (f" #{n + 1}" if len(r["images"]) > 1 else ""), im.resize((int(im.width * s), 520))))
    W = 2200
    rows, row, w = [], [], 0
    for t in tiles:
        if row and w + t[1].width > W:
            rows.append(row); row, w = [], 0
        row.append(t); w += t[1].width + 20
    if row:
        rows.append(row)
    sheet = Image.new("RGB", (W, max(1, len(rows)) * 560 + 10), (240, 240, 240))
    dr = ImageDraw.Draw(sheet)
    y = 5
    for row in rows:
        x = 10
        for name, im in row:
            dr.text((x, y + 4), name, fill=(180, 0, 0)); sheet.paste(im, (x, y + 24)); x += im.width + 20
        y += 560
    out = d / "flagged.jpg"
    sheet.save(out, quality=80)
    return out


CSS = """
:root{--ink:#111827;--muted:#6B7280;--line:#E5E7EB;--bg:#F7F7F8}
*{box-sizing:border-box}body{margin:0;font:15px/1.45 -apple-system,BlinkMacSystemFont,'Segoe UI',Inter,Arial,sans-serif;color:var(--ink);background:var(--bg)}
header{position:sticky;top:0;z-index:5;background:#fff;border-bottom:1px solid var(--line);padding:14px 28px;display:flex;gap:18px;align-items:center;flex-wrap:wrap}
header h1{font-size:18px;margin:0}header .sub{color:var(--muted)}
.filters{display:flex;gap:6px;flex-wrap:wrap;margin-left:auto}.filters button{border:1px solid var(--line);background:#fff;border-radius:999px;padding:6px 12px;cursor:pointer;font:inherit}
.filters button.on{background:var(--ink);color:#fff;border-color:var(--ink)}
main{padding:20px 28px 60px;max-width:1500px;margin:0 auto}
.post{background:#fff;border:1px solid var(--line);border-radius:12px;padding:18px 20px;margin:16px 0}
.post h2{font-size:16px;margin:0 0 2px}.post .meta{color:var(--muted);font-size:13px;margin-bottom:12px}
.opts{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:16px}
.opt{border:1px solid var(--line);border-radius:10px;overflow:hidden;background:#fff}
.opt img{display:block;width:100%;height:auto;background:#ddd}
.opt .b{padding:10px 12px}.opt .k{font-weight:700}.opt .tpl{color:var(--muted);font-size:12px}
.pill{display:inline-block;font-size:12px;border-radius:999px;padding:2px 9px;margin:6px 6px 0 0;background:#F3F4F6}
.pill.Approved{background:#D1FAE5;color:#065F46}.pill.Rejected{background:#FEE2E2;color:#991B1B}
.pill.Changes{background:#FEF3C7;color:#92400E}.pill.fix{background:#FEE2E2;color:#991B1B}.pill.look{background:#FEF3C7;color:#92400E}
details{margin-top:8px;font-size:13px;color:#374151}summary{cursor:pointer;color:var(--muted)}
.strip{display:flex;gap:6px;overflow-x:auto;padding:8px 12px 0}.strip img{height:90px;width:auto;border-radius:4px}
a{color:inherit}
"""


def gallery(c: store.Client, d: Path, rows: list[dict]):
    """index.html: every option, grouped by post, with status and checks. Open it from Drive for desktop."""
    posts = {}
    for r in rows:
        posts.setdefault(r["post"], []).append(r)
    months = sorted({r["month"] for r in rows if r["month"] is not None})
    parts = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width'>"
             f"<title>{html.escape(c.name)} · {html.escape(d.name)}</title><style>{CSS}</style></head><body>"
             f"<header><h1>{html.escape(c.name)}</h1><span class='sub'>{html.escape(d.name)} · {len(posts)} posts · "
             f"{sum(r['status'] == 'Approved' for r in rows)} approved options</span><div class='filters'>"
             f"<button class='on' data-m='all'>All</button>" + "".join(f"<button data-m='{m}'>Month {m}</button>" for m in months if m)
             + "</div></header><main>"]
    for pid, rs in posts.items():
        first = rs[0]
        parts.append(f"<section class='post' data-m='{first['month']}'><h2>{html.escape(pid)} · {html.escape(first['topic'])}</h2>"
                     f"<div class='meta'>{'Month ' + str(first['month']) + ' · week ' + str(first['week']) + ' · ' if first['month'] else ''}"
                     f"{html.escape(first.get('slot_type') or '')} · {html.escape(first.get('format') or '')}</div><div class='opts'>")
        for r in rs:
            st_cls = r["status"].split()[0]
            img = f"<a href='{html.escape(r['html'] or '#')}'><img loading='lazy' src='{html.escape(r['image'] or '')}' alt=''></a>" if r["image"] else "<div style='height:200px;display:grid;place-items:center;color:#999'>not drawn yet</div>"
            strip = ""
            if len(r["images"]) > 1:
                strip = "<div class='strip'>" + "".join(f"<img loading='lazy' src='{html.escape(i)}' alt=''>" for i in r["images"][1:12]) + "</div>"
            pills = f"<span class='pill {st_cls}'>{html.escape(r['status'])}</span>"
            if r["blocking"]:
                pills += f"<span class='pill fix'>Needs fixing</span>"
            elif r["warnings"]:
                pills += f"<span class='pill look'>Look at</span>"
            det = ""
            notes = r["blocking"] + r["warnings"] + [f"Confirm: {x}" for x in r["confirm"]]
            if notes:
                det = "<details><summary>Checks</summary><ul>" + "".join(f"<li>{html.escape(n)}</li>" for n in notes) + "</ul></details>"
            parts.append(f"<div class='opt'>{img}{strip}<div class='b'><span class='k'>Option {html.escape(r['option'])}</span> "
                         f"<span class='tpl'>{html.escape(r['template'])}</span><div>{pills}</div>"
                         f"<details><summary>Caption</summary><p>{html.escape(r['caption'])}</p></details>{det}</div></div>")
        parts.append("</div></section>")
    parts.append("""</main><script>document.querySelectorAll('.filters button').forEach(b=>b.onclick=()=>{
document.querySelectorAll('.filters button').forEach(x=>x.classList.toggle('on',x===b));
document.querySelectorAll('.post').forEach(p=>p.style.display=(b.dataset.m==='all'||p.dataset.m===b.dataset.m)?'':'none');});</script></body></html>""")
    write_text(d / "index.html", "".join(parts))
