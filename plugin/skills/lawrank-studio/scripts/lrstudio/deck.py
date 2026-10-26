"""The client deck (.pptx, opens in PowerPoint and Google Slides): brand board, the designs with their captions,
what the firm needs to confirm, and next steps. Client-facing only: no internal check notes."""
import io
import tempfile
from pathlib import Path

from .util import write_bytes


def _rgb(hexc):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexc.lstrip("#").upper()[:6])


def build_deck(c, d: Path, brand: dict, groups: list[dict], confirm: list[str], out: Path, title: str) -> Path:
    """groups: [{"title", "subtitle", "images": [paths], "captions": [str], "notes": str}] one slide each."""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt
    from PIL import Image
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    ink, accent, paper = brand["colors"]["ink"], brand["colors"]["accent"], brand["colors"].get("paper", "#F5F5F5")
    head = brand["fonts"]["head"]["family"]; body = brand["fonts"]["body"]["family"]
    W, H = prs.slide_width, prs.slide_height
    tmp = Path(tempfile.mkdtemp(prefix="lr-deck-"))

    def small(p):
        """A compressed copy for the deck (decks travel by email; full-size PNGs stay in the posts zip)."""
        im = Image.open(p).convert("RGB")
        if max(im.size) > 1600:
            k = 1600 / max(im.size); im = im.resize((round(im.width * k), round(im.height * k)), Image.LANCZOS)
        q = tmp / (Path(p).stem + ".jpg"); im.save(q, "JPEG", quality=84)
        return q

    def bg(slide, color):
        f = slide.background.fill; f.solid(); f.fore_color.rgb = _rgb(color)

    def text(slide, x, y, w, h, s, size, color, bold=False, font=None):
        tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = s
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = _rgb(color); r.font.name = font or body
        return tb

    def logo_png(variant):
        """The vector logo rasterised for the deck (PowerPoint can't place SVG everywhere)."""
        svg = c.brand_dir / "logos" / f"logo-{variant}.svg"
        if not svg.exists():
            return None
        from playwright.sync_api import sync_playwright
        out_png = tmp / f"logo-{variant}.png"
        with sync_playwright() as p:
            b = p.chromium.launch(); pg = b.new_page(viewport={"width": 900, "height": 300})
            pg.set_content(f"<html><body style='margin:0;background:transparent'><div style='width:880px'>{svg.read_text()}</div>"
                           "<style>svg{width:100%;height:auto;display:block}</style></body></html>")
            pg.locator("svg").first.screenshot(path=str(out_png), omit_background=True); b.close()
        return out_png

    # cover
    s = prs.slides.add_slide(blank); bg(s, ink)
    lw = logo_png("white")
    if lw:
        s.shapes.add_picture(str(lw), Inches(0.8), Inches(0.7), width=Inches(3.2))
    text(s, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.6), title, 44, "#FFFFFF", True, head)
    text(s, Inches(0.8), Inches(4.1), Inches(11.5), Inches(0.8), f"Prepared for {brand['name']} by LawRank", 18, "#D0D4D8")
    bar = s.shapes.add_shape(1, Inches(0.8), Inches(3.95), Inches(1.2), Inches(0.08)); bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(accent); bar.line.fill.background()

    # brand board pages
    for png in sorted((c.brand_dir / "board").glob("board_*.png"))[:4]:
        s = prs.slides.add_slide(blank); bg(s, "#FFFFFF")
        s.shapes.add_picture(str(small(png)), 0, 0, width=W)

    # designs
    for g in groups:
        s = prs.slides.add_slide(blank); bg(s, "#F4F5F6")
        text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6), g["title"], 24, ink, True, head)
        if g.get("subtitle"):
            text(s, Inches(0.6), Inches(0.9), Inches(12), Inches(0.4), g["subtitle"], 13, "#5B6470")
        imgs = [p for p in g["images"] if Path(p).exists()][:4]
        if imgs:
            gap, top, avail_w, avail_h = Inches(0.3), Inches(1.5), W - Inches(1.2), Inches(4.4)
            sizes = [Image.open(p).size for p in imgs]
            hgt = min(avail_h, int((avail_w - gap * (len(imgs) - 1)) / sum(w / h for w, h in sizes)))
            x = Inches(0.6)
            for p, (w, h) in zip(imgs, sizes):
                wd = int(hgt * w / h)
                s.shapes.add_picture(str(small(p)), x, top, width=wd, height=hgt); x += wd + gap
        caps = [cp for cp in g.get("captions", []) if cp]
        if caps:
            text(s, Inches(0.6), Inches(6.05), Inches(12.1), Inches(1.2), "Caption: " + caps[0], 12, "#3A434D")

    # confirm + next steps
    s = prs.slides.add_slide(blank); bg(s, "#FFFFFF")
    text(s, Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.8), "Please confirm", 32, ink, True, head)
    items = confirm or ["Nothing to confirm in this set."]
    tb = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items[:14]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = "•  " + it; r.font.size = Pt(15); r.font.name = body; r.font.color.rgb = _rgb("#2B333C")
    s = prs.slides.add_slide(blank); bg(s, ink)
    text(s, Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.8), "Next steps", 32, "#FFFFFF", True, head)
    text(s, Inches(0.8), Inches(1.7), Inches(11.5), Inches(4),
         "1. Review the designs and captions in this deck.\n2. Confirm the items on the previous page.\n"
         "3. Tell us which options to post (or what to change).\nWe'll prepare the final files and schedule.", 18, "#E6E9EC")
    buf = io.BytesIO(); prs.save(buf)
    write_bytes(out, buf.getvalue())
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)
    return out
