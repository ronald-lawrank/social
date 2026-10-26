"""Clients on disk (in the shared Drive folder).

Clients/<Firm name>/
  client.json        identity, practice area, package, chosen direction
  brand/brand.json   visual identity + voice (studio schema, see brand.py)
  brand/facts.xlsx   facts with source and approval columns (people edit this in Google Sheets)
  brand/photos/, brand/logos/, brand/site/, brand/board.html
  campaigns/<name>/  posts/, review.xlsx, previews/, images/, video/, exports/
  history.log
"""
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from . import paths
from .util import LRError, now, read_json, safe_folder, slugify, write_json


def domain_of(url: str) -> str:
    u = url if re.match(r"^https?://", url or "", re.I) else "https://" + (url or "")
    host = (urlparse(u).hostname or "").lower()
    return host[4:] if host.startswith("www.") else host


def normalize_url(url: str) -> str:
    url = (url or "").strip()
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    p = urlparse(url)
    if not p.hostname or "." not in p.hostname:
        raise LRError(f"'{url}' doesn't look like a website address", fix="Give the firm's site, e.g. smithinjury.com")
    host = p.hostname.lower() + (f":{p.port}" if p.port else "")     # keep a port; drop any user:password@
    return f"{p.scheme.lower()}://{host}{p.path or '/'}"


_LEAD = re.compile(r"^(?:the\s+)?(?:(?:law\s+)?offices?\s+of|law\s+firm\s+of|firm\s+of)\s+|^the\s+", re.I)
_TAIL = re.compile(r",.*$|\s+&\s+.*$|\s+(?:law|attorneys?|lawyers?|legal|injury|trial|firm|group)\b.*$|"
                   r"\s+(?:pllc|llp|llc|pc|p\.c\.|pa|p\.a\.|ltd|inc)\.?$", re.I)


def short_name(name: str) -> str:
    """The name the firm goes by in copy: "The Law Office of Ben Carrasco, PLLC" -> "Ben Carrasco",
    "Smith & Jones Law Firm" -> "Smith", "The McMinn Law Firm" -> "McMinn". Never a bare article."""
    n = re.sub(r"\s+", " ", name or "").strip()
    s = _LEAD.sub("", n).strip()
    s = _TAIL.sub("", s).strip(" ,.-") or s
    if not s or s.lower() in ("the", "a", "an", "law", "office", "offices", "firm"):
        s = n
    return s[:24].strip()


@dataclass
class Client:
    dir: Path
    data: dict

    @property
    def name(self) -> str: return self.data["name"]
    @property
    def slug(self) -> str: return self.data["slug"]
    @property
    def brand_dir(self) -> Path: return self.dir / "brand"
    @property
    def campaigns_dir(self) -> Path: return self.dir / "campaigns"

    def brand(self) -> dict:
        b = read_json(self.brand_dir / "brand.json")
        if not b:
            raise LRError(f"{self.name} has no brand profile yet", fix=f"Run: lr client new {self.data.get('website')}")
        return b

    def save(self):
        self.data["updated"] = now()
        write_json(self.dir / "client.json", self.data)


def all_clients() -> list[Client]:
    root = paths.clients_dir()
    out = []
    for d in sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")):
        data = read_json(d / "client.json")
        if data:
            out.append(Client(d, data))
    return out


def find(query: str) -> Client:
    """By slug, folder name, firm name or website domain; case-insensitive, partial match if unique."""
    q = query.strip().lower()
    qd = domain_of(q) if "." in q else None
    cands = all_clients()
    for test in (lambda c: c.slug == slugify(q) or c.dir.name.lower() == q or c.name.lower() == q,
                 lambda c: qd and c.data.get("domain") == qd,
                 lambda c: q in c.name.lower() or q in c.slug or slugify(q) in c.slug):
        hits = [c for c in cands if test(c)]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            raise LRError(f"'{query}' matches several clients: {', '.join(c.name for c in hits)}", fix="Use the full firm name.")
    raise LRError(f"No client matches '{query}'", fix="Run `lr client list`, or `lr client new <website>` for a new firm.",
                  clients=[c.name for c in cands][:30])


def find_by_domain(domain: str) -> Client | None:
    for c in all_clients():
        if c.data.get("domain") == domain:
            return c
    return None


def create(name: str, website: str, **fields) -> Client:
    url = normalize_url(website)
    dom = domain_of(url)
    if find_by_domain(dom):
        raise LRError(f"{dom} is already a client", fix=f"Use `lr client show {dom}`.")
    folder = paths.clients_dir() / safe_folder(name)
    n = 2
    while folder.exists():
        folder = paths.clients_dir() / f"{safe_folder(name)} ({n})"; n += 1
    taken = {c.slug for c in all_clients()}
    slug, n = slugify(fields.pop("short", None) or name), 2
    base = slug
    while slug in taken:
        slug, n = f"{base}-{n}", n + 1
    data = {"name": name, "slug": slug, "website": url, "domain": dom, "practice_area": None, "package": "growth",
            "direction": None, "status": "onboarding", "created": now(), **fields}
    for sub in ("brand/logos", "brand/photos", "brand/site", "campaigns"):
        (folder / sub).mkdir(parents=True, exist_ok=True)
    c = Client(folder, data); c.save()
    return c


# ---------------------------------------------------------------- facts.xlsx
FACT_COLS = ["Key", "Fact", "Category", "Source page", "Excerpt", "Found on site", "Needs approval", "Approved",
             "Approved by", "Notes"]
YES = {"yes", "y", "true", "1", "approved", "ok", "✓", "x"}


def _yes(v) -> bool:
    return str(v or "").strip().lower() in YES


def read_facts(client: Client) -> dict[str, dict]:
    """facts.xlsx -> {key: {...}}. People edit this file in Google Sheets; it's the record of approvals."""
    f = client.brand_dir / "facts.xlsx"
    if not f.exists():
        return {}
    from openpyxl import load_workbook
    try:
        ws = load_workbook(f, read_only=True, data_only=True).worksheets[0]
    except Exception as e:
        raise LRError(f"Can't read facts.xlsx ({e})", fix="Open it in Google Sheets and check it isn't damaged; Drive keeps earlier versions.")
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    idx = {h: i for i, h in enumerate(header)}
    out = {}
    for r in rows:
        get = lambda col: (r[idx[col]] if col in idx and idx[col] < len(r) else None)
        key = str(get("Key") or "").strip()
        if not key:
            continue
        needs = _yes(get("Needs approval"))
        approved = _yes(get("Approved"))
        found = _yes(get("Found on site"))
        out[key] = {"key": key, "text": str(get("Fact") or "").strip(), "category": str(get("Category") or "general"),
                    "source": get("Source page"), "excerpt": get("Excerpt"), "verified": found,
                    "needs_approval": needs, "approved": approved, "approved_by": get("Approved by"), "notes": get("Notes"),
                    "usable": approved or (found and not needs),        # a person's approval counts even for facts not on the site
                    "status": "usable" if approved or (found and not needs) else ("needs approval" if found else "not verified")}
    return out


def write_facts(client: Client, facts: list[dict], replace: bool = False):
    """Merge facts into facts.xlsx. Existing rows keep every human edit (text, approval, notes); new keys are added."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    existing = {} if replace else read_facts(client)
    rows = []
    for k, f in existing.items():
        rows.append([k, f["text"], f["category"], f["source"], f["excerpt"], "Yes" if f["verified"] else "No",
                     "Yes" if f["needs_approval"] else "No", "Yes" if f["approved"] else "", f["approved_by"] or "", f["notes"] or ""])
    for f in facts:
        if f["key"] in existing:
            continue
        rows.append([f["key"], f["text"], f.get("category", "general"), f.get("source", ""), f.get("excerpt", ""),
                     "Yes" if f.get("verified") else "No", "Yes" if f.get("needs_approval") else "No", "", "", f.get("notes") or ""])
    wb = Workbook(); ws = wb.active; ws.title = "Facts"
    ws.append(FACT_COLS)
    for r in rows:
        ws.append(r)
    head = PatternFill("solid", fgColor="1F2937")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head
    widths = [18, 70, 14, 40, 40, 13, 15, 11, 16, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[6].value == "Yes" and row[7].value != "Yes":
            row[7].fill = PatternFill("solid", fgColor="FDE68A")       # waiting for sign-off
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"H2:H{max(2, ws.max_row + 200)}")
    ws.freeze_panes = "B2"
    from .util import write_bytes
    import io
    buf = io.BytesIO(); wb.save(buf)
    write_bytes(client.brand_dir / "facts.xlsx", buf.getvalue())
